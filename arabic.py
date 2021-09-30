from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.table import WD_TABLE_DIRECTION
import pyarabic.araby as araby
import pyarabic.number as number

document = Document()

text = u'ارض زراعية مسقية'
text2 = u'ارض مسقية مغروسة'

table2 = document.add_table(rows=1, cols=12, style='Table Grid')
# table2.alignment = WD_TABLE_ALIGNMENT.RIGHT
table2.alignment = WD_TABLE_DIRECTION.RTL
table2.rows.style = "borderColor:red"

table2.cell(0, 2).text = text
table2.cell(0, 3).text = text2

document.save('test.docx')